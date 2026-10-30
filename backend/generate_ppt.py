"""
Generate an 8-slide widescreen (16:9) hackathon pitch deck (.pptx)
for Tsukuyomi Watch: The Samurai's Celestial Guide (PECathon 2026 - Team PF-XBZYZH).
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Theme Palette (Sumi-e Obsidian, Vermilion Crimson, Samurai Gold, Starlight Ivory)
BG_DARK = RGBColor(5, 7, 10)
CARD_BG = RGBColor(12, 16, 24)
CARD_ALT = RGBColor(18, 24, 36)
GOLD = RGBColor(201, 162, 74)
GOLD_BRIGHT = RGBColor(252, 227, 138)
CRIMSON = RGBColor(224, 35, 28)
TEXT_MAIN = RGBColor(223, 231, 224)
TEXT_MUTED = RGBColor(139, 150, 143)
EMERALD = RGBColor(52, 211, 153)
WHITE = RGBColor(255, 255, 255)


def set_slide_background(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK

    # Top crimson-gold accent bar
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = CRIMSON
    bar.line.fill.background()


def add_header(slide, slide_num: int, kicker: str, title: str, subtitle: str):
    set_slide_background(slide)

    # Kicker
    tb_k = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(9.5), Inches(0.35))
    tf_k = tb_k.text_frame
    tf_k.word_wrap = True
    p_k = tf_k.paragraphs[0]
    p_k.text = kicker.upper()
    p_k.font.size = Pt(10)
    p_k.font.bold = True
    p_k.font.color.rgb = CRIMSON
    p_k.font.name = "Consolas"

    # Slide Number Badge (top right)
    tb_n = slide.shapes.add_textbox(Inches(10.2), Inches(0.32), Inches(2.5), Inches(0.35))
    tf_n = tb_n.text_frame
    p_n = tf_n.paragraphs[0]
    p_n.alignment = PP_ALIGN.RIGHT
    p_n.text = f"PECathon 2026 • Slide {slide_num} / 8"
    p_n.font.size = Pt(10)
    p_n.font.bold = True
    p_n.font.color.rgb = GOLD
    p_n.font.name = "Consolas"

    # Main Title
    tb_t = slide.shapes.add_textbox(Inches(0.7), Inches(0.68), Inches(11.9), Inches(0.65))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = WHITE
    p_t.font.name = "Georgia"

    # Subtitle
    tb_s = slide.shapes.add_textbox(Inches(0.7), Inches(1.32), Inches(11.9), Inches(0.55))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True
    p_s = tf_s.paragraphs[0]
    p_s.text = subtitle
    p_s.font.size = Pt(13)
    p_s.font.color.rgb = TEXT_MUTED
    p_s.font.name = "Calibri"


def add_card(slide, left: float, top: float, width: float, height: float, title: str, bullets: list[str], border_color=GOLD):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.25)

    # Top inner strip
    strip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left + 0.15), Inches(top + 0.08), Inches(width - 0.3), Inches(0.04)
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = CRIMSON
    strip.line.fill.background()

    tb = slide.shapes.add_textbox(
        Inches(left + 0.22), Inches(top + 0.2), Inches(width - 0.44), Inches(height - 0.3)
    )
    tf = tb.text_frame
    tf.word_wrap = True

    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(15)
    p_title.font.bold = True
    p_title.font.color.rgb = GOLD_BRIGHT
    p_title.font.name = "Georgia"
    p_title.space_after = Pt(8)

    for b in bullets:
        p = tf.add_paragraph()
        p.text = f"✦  {b}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_MAIN
        p.font.name = "Calibri"
        p.space_after = Pt(5)


def build_deck(output_path: str):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TITLE & VISION
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(2.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "PECATHON 2026 OFFICIAL SUBMISSION  •  TEAM PF-XBZYZH  •  月読命"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = CRIMSON
    p0.font.name = "Consolas"
    p0.space_after = Pt(10)

    p1 = tf1.add_paragraph()
    p1.text = "TSUKUYOMI WATCH"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.font.name = "Georgia"

    p2 = tf1.add_paragraph()
    p2.text = "The Samurai's Celestial Guide — Scientific Astrophotography, Photutils Source Detection & Live Ephemeris Sanctuary"
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = GOLD_BRIGHT
    p2.font.name = "Georgia"
    p2.space_after = Pt(8)

    # 4 Pillars on Slide 1
    add_card(
        s1, 0.8, 3.1, 2.78, 2.95,
        "I. Watchtower Map",
        [
            "Interactive Leaflet dark-sky sanctuary map.",
            "Overpass API observatory & viewpoint finder.",
            "Transparent Night Honor Score formula.",
        ],
    )
    add_card(
        s1, 3.78, 3.1, 2.78, 2.95,
        "II. Spirit Lens AI",
        [
            "Real-ESRGAN + OpenCV CLAHE restoration.",
            "Photutils & Astropy star/galaxy detection.",
            "Original vs. Restored artifact cross-matching.",
        ],
    )
    add_card(
        s1, 6.76, 3.1, 2.78, 2.95,
        "III. Live Ephemeris",
        [
            "US Naval Observatory (USNO) Moon Phases.",
            "IMCCE OPALE Lunar Eclipse predictions.",
            "Open-Meteo hourly cloud, humidity & visibility.",
        ],
    )
    add_card(
        s1, 9.74, 3.1, 2.78, 2.95,
        "IV. Samurai Honours",
        [
            "Auth-gated Chapter V dedicated web page.",
            "88 IAU Constellations + EXIF camera audit.",
            "Supabase Cloud Postgres & Storage bucket.",
        ],
    )

    # Footer bar on Slide 1
    tb_f = s1.shapes.add_textbox(Inches(0.8), Inches(6.35), Inches(11.7), Inches(0.6))
    tf_f = tb_f.text_frame
    pf = tf_f.paragraphs[0]
    pf.text = "LIVE PRODUCTION: https://tsukuyomi-watch-mu.vercel.app   |   GITHUB: github.com/PECathon-2026/Team-PF-XBZYZH   |   41/41 PYTEST PASSING"
    pf.font.size = Pt(11)
    pf.font.bold = True
    pf.font.color.rgb = EMERALD
    pf.font.name = "Consolas"

    # =========================================================================
    # SLIDE 2: THE PROBLEM LANDSCAPE
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(
        s2, 2,
        "01 • Problem Statement",
        "Why Current Stargazing & Astrophotography Tools Fail",
        "Amateur astronomers, students, and citizen scientists face four critical pain points between sky observation and scientific truth.",
    )
    add_card(
        s2, 0.7, 2.05, 5.75, 2.3,
        "1. Generative AI Hallucinations in Night Photos",
        [
            "Consumer AI enhancers fabricate non-existent stars and nebulae.",
            "Destroys photometric integrity and misleads amateur observers.",
            "Our Solution: Deterministic CLAHE + Photutils cross-matching that flags 'restored_only' artifacts.",
        ],
    )
    add_card(
        s2, 6.85, 2.05, 5.75, 2.3,
        "2. Opaque Black-Box Stargazing Scores",
        [
            "Existing apps output arbitrary visibility scores without disclosing formulas.",
            "Conflate numerical weather forecasts (NWP) with live telescope feeds.",
            "Our Solution: Disclosed mathematical weights + automatic weight rebalancing when sensors are missing.",
        ],
    )
    add_card(
        s2, 0.7, 4.6, 5.75, 2.3,
        "3. Fragmented Ephemeris & Archival Portals",
        [
            "Checking USNO moon phases, IMCCE eclipses, and NASA SkyView plates requires 6+ legacy tools.",
            "Timestamps are rarely normalized from UTC to the observer's local timezone.",
            "Our Solution: Unified sanctuary console with live USNO, IMCCE, Open-Meteo & NASA SkyView drawers.",
        ],
    )
    add_card(
        s2, 6.85, 4.6, 5.75, 2.3,
        "4. Unverified Community Observation Leaderboards",
        [
            "Social astronomy walls are flooded with Midjourney/Canva edits and stripped screenshots.",
            "Our Solution: 5-step EXIF hardware sensor, GPS coordinate, and night-window (17:00-06:00) verification stored in Supabase.",
        ],
    )

    # =========================================================================
    # SLIDE 3: SYSTEM ARCHITECTURE
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(
        s3, 3,
        "02 • System Architecture",
        "Full-Stack Scientific & Cloud Architecture",
        "Engineered for zero-downtime resilience: React 19 + WebGL frontend on Vercel, FastAPI scientific engine, and Supabase Cloud Postgres.",
    )
    add_card(
        s3, 0.7, 2.05, 3.8, 4.8,
        "Frontend Sanctuary (Vercel)",
        [
            "React 19 + TypeScript + Vite 8 + Tailwind CSS v4.",
            "Section 0: Fullscreen 3D Kage Temple WebGL scene.",
            "5-Button Navbar: Watchtower, Celestial Scroll, Spirit Lens, Dojo Score, and Auth-Gated Samurai Honours.",
            "Client-side Canvas CLAHE & 2D MAD source detection fallback guarantees 0ms downtime.",
        ],
    )
    add_card(
        s3, 4.75, 2.05, 3.8, 4.8,
        "Scientific Backend (FastAPI)",
        [
            "Python 3.11 + FastAPI + Pydantic validation.",
            "Astropy & Photutils: Background2D, DAOStarFinder, detect_sources, deblend_sources, SourceCatalog.",
            "Real-ESRGAN (SRVGGNetCompact & RRDBNet) + OpenCV restoration.",
            "Multi-wavelength NASA SkyView & Transient Spectral routers.",
        ],
    )
    add_card(
        s3, 8.8, 2.05, 3.8, 4.8,
        "Supabase Cloud + Public APIs",
        [
            "Supabase Postgres: public.stargazer_users & public.celestial_uploads tables with RLS.",
            "Supabase Storage: Public 'celestial-uploads' bucket for verified sky exposures.",
            "Live USNO Moon Phases, IMCCE OPALE Lunar Eclipses, and Open-Meteo NWP Forecast APIs.",
        ],
    )

    # =========================================================================
    # SLIDE 4: SPIRIT LENS & PHOTUTILS DETECTION
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(
        s4, 4,
        "03 • Chapter III: Spirit Lens",
        "Real-ESRGAN Restoration & Photutils Candidate Detection",
        "Separates non-destructive display stretches (linear, percentile, log, asinh) from scientific floating-point ADU photometry.",
    )
    add_card(
        s4, 0.7, 2.05, 5.8, 4.8,
        "5-Stage Astronomical Detection Pipeline",
        [
            "1. Image & FITS Validation: Decodes JPEG, PNG, WebP, TIFF & FITS HDUs (astropy.io.fits); masks RGBA alpha channels.",
            "2. 2D Background Modeling: Photutils Background2D + MedianBackground + SigmaClip(3.0) across adaptive meshes.",
            "3. Stellar Point-Source Pass: DAOStarFinder + CircularAperture photometry for unresolved stars.",
            "4. Extended Galaxy Segmentation: 2D Gaussian PSF filter + detect_sources + deblend_sources + SourceCatalog.",
            "5. Provenance Cross-Match: Matches original vs. restored detections ('both', 'original_only', 'restored_only').",
        ],
    )
    add_card(
        s4, 6.8, 2.05, 5.8, 4.8,
        "Measured Properties & Scientific Honesty",
        [
            "Astrometry: Zero-based pixel centroid (x_centroid, y_centroid) & bounding box.",
            "Photometry: Peak ADU, background-subtracted segment flux, circular aperture flux, local RMS & SNR.",
            "Morphology: Radial FWHM (px), area (px²), semi-major/minor axes (a, b), ellipticity (1 - b/a), orientation (deg).",
            "4 Candidate Classes: Star-Like Point, Extended-Source, Possible Galaxy, or Uncertain/Artifact.",
            "Interactive HTML5 Canvas overlay + zoomed cutout inspector + CSV & JSON export.",
        ],
    )

    # =========================================================================
    # SLIDE 5: LIVE PUBLIC APIS & NIGHT HONOR SCORE
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(
        s5, 5,
        "04 • Chapters I, II & IV: Live Telemetry",
        "USNO, IMCCE & Open-Meteo Integration + Night Honor Score",
        "Replaces static mock data with live scientific ephemerides and weather forecasts, complete with local timezone conversion and TTL caching.",
    )
    add_card(
        s5, 0.7, 2.05, 3.8, 3.1,
        "1. US Naval Observatory (USNO)",
        [
            "Endpoint: aa.usno.navy.mil/api/moon/phases/year",
            "Fetches official New Moon, First Quarter, Full Moon & Last Quarter UTC timestamps.",
            "Computes instantaneous illumination % and lunar age (days).",
        ],
    )
    add_card(
        s5, 4.75, 2.05, 3.8, 3.1,
        "2. IMCCE OPALE Lunar Eclipses",
        [
            "Endpoint: opale.imcce.fr/api/v1/phenomena/eclipses/301",
            "INPOP19A ephemeris for Total, Partial & Penumbral lunar eclipses.",
            "Extracts umbral magnitude, durations & zenith coordinates.",
        ],
    )
    add_card(
        s5, 8.8, 2.05, 3.8, 3.1,
        "3. Open-Meteo Weather API",
        [
            "Endpoint: api.open-meteo.com/v1/forecast",
            "Dynamic coordinate hourly cloud cover, 2m humidity & optical visibility.",
            "Explicitly badged as NWP Forecast (Never Live Telescope Data).",
        ],
    )
    add_card(
        s5, 0.7, 5.35, 11.9, 1.55,
        "Transparent Night Honor Score Mathematical Formula",
        [
            "Night Honor Score = (0.53 × ClearSky) + (0.29 × VisibilityScore) + (0.18 × HumidityScore)   |   Automatically rebalances weights if light-pollution sensor data is absent.",
        ],
        border_color=CRIMSON,
    )

    # =========================================================================
    # SLIDE 6: CHAPTER V — SAMURAI HONOURS & SUPABASE DATABASE
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(
        s6, 6,
        "05 • Chapter V: Samurai Honours & Supabase Cloud",
        "Auth-Gated Samurai Honours Web Page & Database Schema",
        "Clicking 'Samurai Honours' in the Navbar triggers Stargazer authentication before unlocking the dedicated Chapter V web page.",
    )
    add_card(
        s6, 0.7, 2.05, 5.8, 4.8,
        "Table 1: public.stargazer_users (Authentication)",
        [
            "id (uuid PK) & auth_user_id (FK -> auth.users)",
            "email (text UNIQUE NOT NULL) & name (text NOT NULL)",
            "password_hash (SHA-256 salted passphrase digest)",
            "location (text) — Home dark-sky sanctuary ground",
            "avatar (text) — Samurai Clan Crest (🌸, 🌙, 🦉, 🏮, 🔭, ⭐)",
            "honor_xp (integer) & rank_title ('Star Seeker' -> 'Grand Astronomer')",
            "stamps (text[]) & registered_at / last_login_at (timestamptz)",
        ],
    )
    add_card(
        s6, 6.8, 2.05, 5.8, 4.8,
        "Table 2: public.celestial_uploads (Images & EXIF)",
        [
            "id (uuid PK) & user_id (FK -> public.stargazer_users.id)",
            "item_id, item_name & category (88 Constellations, Planets, Deep Sky)",
            "camera_make, camera_model & camera_gear (EXIF hardware check)",
            "captured_at (timestamptz) — Night window check (17:00–06:00)",
            "gps_latitude & gps_longitude (double precision dark-sky coords)",
            "image_url & storage_path (Supabase Storage bucket 'celestial-uploads')",
            "verification_passed, match_score, honor_awarded & exif_checks (jsonb)",
        ],
    )

    # =========================================================================
    # SLIDE 7: SIDECAR OBSERVATORIES (NASA SKYVIEW & TRANSIENT LAB)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(
        s7, 7,
        "06 • Multi-Spectral & Time-Domain Sidecar Labs",
        "NASA SkyView Virtual Observatory & Transient Spectral Lab",
        "Slide-out observatory consoles accessible from any chapter for multi-wavelength survey comparison and supernova spectroscopy.",
    )
    add_card(
        s7, 0.7, 2.05, 5.8, 4.8,
        "Sky Lantern Lens — NASA GSFC SkyView",
        [
            "Queries 5 archival wavelength bands: Optical (DSS2 Red), Near-IR (2MASS), Mid-IR (WISE 3.4μm), UV (GALEX), and X-Ray (ROSAT).",
            "Live Simbad / Sesame astronomical name resolver (M31, M42, M45, NGC 1068 -> J2000 RA/Dec).",
            "Interactive side-by-side reticle comparator to inspect user astrophotography against archival NASA HEASARC plates.",
            "Awards Samurai Honor XP when observations are logged to the personal catalog.",
        ],
    )
    add_card(
        s7, 6.8, 2.05, 5.8, 4.8,
        "Transient Spectral & Time-Domain Divination Lab",
        [
            "Classifies transient light curves & spectra across Supernovae (SN Ia, SN II), Kilonovae, Tidal Disruption Events (TDE), and AGN.",
            "Relativistic Doppler shift calculator computes ejecta expansion velocity (km/s and % speed of light) from rest vs. observed lines.",
            "Built-in Web Audio API astronomical sonifier translates magnitude light curves into real-time pitch and timbre sweeps.",
        ],
    )

    # =========================================================================
    # SLIDE 8: VERIFICATION, DEPLOYMENT & CONCLUSION
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(
        s8, 8,
        "07 • Verification, Cloud Deployment & Impact",
        "Production-Ready, Scientifically Verified & Live",
        "Complete end-to-end engineering validated by 41 automated backend tests and deployed globally on Vercel + Supabase Cloud.",
    )
    add_card(
        s8, 0.7, 2.05, 3.8, 3.0,
        "41 / 41 Automated Tests",
        [
            "35 Photutils/Astropy detection & FITS validation tests.",
            "3 NASA SkyView & Transient lab API tests.",
            "3 Live USNO, IMCCE & Open-Meteo integration tests.",
        ],
    )
    add_card(
        s8, 4.75, 2.05, 3.8, 3.0,
        "Live Cloud Infrastructure",
        [
            "Frontend: Deployed on Vercel Global Edge CDN.",
            "Database & Storage: Live on Supabase Postgres (lcwsuzobacqzbutqwmvh).",
            "Backend: Containerized via Dockerfile & render.yaml.",
        ],
    )
    add_card(
        s8, 8.8, 2.05, 3.8, 3.0,
        "Key Hackathon Differentiators",
        [
            "Real scientific Astropy/Photutils code — zero fake overlays.",
            "Anti-GenAI EXIF camera & GPS verification game.",
            "Unified Japanese Samurai celestial UX.",
        ],
    )
    add_card(
        s8, 0.7, 5.3, 11.9, 1.6,
        "Live Links for Judges & Reviewers",
        [
            "Production Web App: https://tsukuyomi-watch-mu.vercel.app   |   Official GitHub Repo: https://github.com/PECathon-2026/Team-PF-XBZYZH",
        ],
        border_color=EMERALD,
    )

    prs.save(output_path)
    print(f"Saved 8-slide PowerPoint presentation to: {output_path}")


if __name__ == "__main__":
    import os
    out_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    build_deck(os.path.join(out_dir, "Tsukuyomi_Watch_PECathon_2026_Pitch_Deck.pptx"))
    build_deck(os.path.join(out_dir, "public", "Tsukuyomi_Watch_PECathon_2026_Pitch_Deck.pptx"))
